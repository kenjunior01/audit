"""
Testes do Excel Copiloto IA (assistente NL→Excel).
Cobre: deteção de intenções, parsing de números PT/EN, executor de
operações, whitelist de segurança, pack de fórmulas PT/EN, endpoints
/excel/assistant e /excel/assistant/apply.
"""
import io
import json

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from dashboard.models import ApiToken
from dashboard import excel_ai_service as ai


def make_xlsx(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Folha1")
    return buf.getvalue()


SAMPLE = pd.DataFrame([
    {"transaction_id": "NF-1", "vendor": "Alpha", "amount": 1200.0,
     "currency": "EUR", "timestamp": "2025-03-04", "category": "TI",
     "user_id": "u1", "status": "Pending"},
    {"transaction_id": "NF-2", "vendor": "Beta", "amount": 25000.0,
     "currency": "EUR", "timestamp": "2025-03-05", "category": "Admin",
     "user_id": "u2", "status": "Approved"},
    {"transaction_id": "NF-3", "vendor": "Alpha", "amount": 900.5,
     "currency": "EUR", "timestamp": "2025-04-10", "category": "TI",
     "user_id": "u1", "status": "Pending"},
])


class CopilotUnitTest(TestCase):
    def test_detect_intents(self):
        self.assertEqual(ai.detect_intents("sugere fórmulas para automatizar")[0], "formula")
        self.assertEqual(ai.detect_intents("mostra o top 5 fornecedores")[0], "top_n")
        self.assertEqual(ai.detect_intents("filtra valores maior que 10000")[0], "filter")
        self.assertEqual(ai.detect_intents("adiciona coluna do mês")[0], "add_column")
        self.assertEqual(ai.detect_intents("tem duplicados aqui?")[0], "duplicates")
        self.assertEqual(ai.detect_intents("como fazer um dashboard?")[0], "dashboard")
        self.assertEqual(ai.detect_intents("um gráfico de barras")[0], "chart")

    def test_parse_number_pt_en(self):
        self.assertEqual(ai._parse_number("10000"), 10000.0)
        self.assertEqual(ai._parse_number("10.000"), 10000.0)
        self.assertEqual(ai._parse_number("10,000"), 10000.0)
        self.assertEqual(ai._parse_number("10k"), 10000.0)
        self.assertIsNone(ai._parse_number("abc"))

    def test_suggest_formulas_semantics(self):
        pack = ai.suggest_formulas(SAMPLE)
        ids = [f["id"] for f in pack]
        for expected in ("kpi_total", "top10_fornecedores", "col_mes",
                         "dupe_id", "procx", "tabela"):
            self.assertIn(expected, ids)
        total = next(f for f in pack if f["id"] == "kpi_total")
        self.assertTrue(total["formula_pt"].startswith("=SOMA("))
        self.assertTrue(total["formula_en"].startswith("=SUM("))
        # fórmulas multi-argumento usam ';' em PT e ',' em EN
        somases = next(f for f in pack if f["id"] == "kpi_estado")
        self.assertIn(";", somases["formula_pt"])
        self.assertNotIn(";", somases["formula_en"])

    def test_suggest_formulas_use_original_headers(self):
        headers = {"amount": "Valor da Nota", "vendor": "Fornecedor Real"}
        pack = ai.suggest_formulas(SAMPLE, headers)
        total = next(f for f in pack if f["id"] == "kpi_total")
        self.assertIn("tblDados[Valor da Nota]", total["formula_pt"])
        self.assertIn("tblDados[Valor da Nota]", total["formula_en"])

    def test_apply_operations_suite(self):
        ops = [
            {"op": "filter", "column": "amount", "gt": 1000},
            {"op": "sort", "by": "amount", "ascending": False},
            {"op": "add_column", "kind": "pct_of_total", "name": "% do Total"},
        ]
        df2, log = ai.apply_operations(SAMPLE, ops)
        self.assertEqual(len(df2), 2)                      # 1200 e 25000
        self.assertEqual(df2.iloc[0]["amount"], 25000.0)   # desc
        self.assertIn("% do Total", df2.columns)
        self.assertAlmostEqual(df2["% do Total"].sum(), 1.0, places=6)

        agg, log2 = ai.apply_operations(SAMPLE, [
            {"op": "aggregate", "group_by": "vendor", "metric": "amount",
             "agg": "sum", "top": 5}])
        self.assertEqual(list(agg.columns), ["vendor", "Total"])
        self.assertEqual(len(agg), 2)
        self.assertEqual(agg.iloc[0]["vendor"], "Beta")   # 25000 > 2100.5

        month, _ = ai.apply_operations(SAMPLE, [
            {"op": "add_column", "kind": "month", "name": "Mês"}])
        self.assertIn("Mês", month.columns)
        self.assertTrue(set(month["Mês"].dropna()) <= {"2025-03", "2025-04"})

        dedup, logd = ai.apply_operations(pd.concat([SAMPLE, SAMPLE.iloc[[0]]]),
                                          [{"op": "remove_duplicates",
                                            "subset": ["transaction_id"]}])
        self.assertEqual(len(dedup), 3)

    def test_validate_ops_whitelist(self):
        ops = [
            {"op": "DROP TABLE users"},                      # inválida
            {"op": "sort", "by": "amount", "ascending": False},
            {"op": "filter", "column": "vendor", "contains": "x'; --"},
            {"op": "top_n", "by": "amount", "n": "15"},
        ]
        clean = ai._validate_ops(ops, SAMPLE)
        self.assertEqual([c["op"] for c in clean], ["sort", "filter", "top_n"])
        self.assertEqual(clean[2]["n"], 15)

    def test_rule_plan_formula_intent(self):
        ctx = ai.describe_frame(SAMPLE)
        plan = ai.rule_plan("sugere fórmulas para o ficheiro ficar automático", ctx)
        self.assertEqual(plan["intent"], "formula")
        self.assertIsNone(plan["formula_ids"])   # pack completo

    def test_rule_plan_top_n(self):
        ctx = ai.describe_frame(SAMPLE)
        plan = ai.rule_plan("top 2 fornecedores por valor", ctx)
        self.assertEqual(plan["intent"], "top_n")
        self.assertTrue(any(o["op"] == "aggregate" for o in plan["operations"]))


class CopilotRESTTest(TestCase):
    def setUp(self):
        self.token = "tok-copilot-1"
        ApiToken.objects.create(token=self.token, user_id="test_admin", role="admin")
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.token)
        self.file = SimpleUploadedFile(
            "dados.xlsx", make_xlsx(SAMPLE.rename(columns={
                "amount": "Valor", "vendor": "Fornecedor",
                "timestamp": "Data", "transaction_id": "Documento",
                "category": "Categoria", "user_id": "Usuário",
                "status": "Estado", "currency": "Moeda"})),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def _post(self, question, **extra):
        payload = {"file": self.file, "question": question}
        payload.update(extra)
        return self.client.post("/django/api/excel/assistant", payload, format="multipart")

    def test_assistant_aggregate(self):
        r = self._post("qual o total por fornecedor?")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        data = r.json()
        self.assertIn(data["intent"], ("aggregate", "top_n"))
        self.assertGreater(len(data["result"]["rows"]), 0)
        self.assertGreater(len(data["formulas"]), 0)
        self.assertIn("columns", data["result"])

    def test_assistant_formula_pack(self):
        r = self._post("sugere fórmulas para tornar o ficheiro mais automático")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        data = r.json()
        self.assertGreaterEqual(len(data["formulas"]), 5)
        for f in data["formulas"]:
            self.assertIn("formula_pt", f)
            self.assertIn("formula_en", f)
        total = next(f for f in data["formulas"] if f["id"] == "kpi_total")
        self.assertTrue(total["formula_en"].startswith("=SUM(tblDados[Valor]"))

    def test_assistant_filter(self):
        r = self._post("filtra valores maior que 5000")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        data = r.json()
        self.assertEqual(data["result"]["total_rows"], 1)

    def test_assistant_requires_question(self):
        r = self.client.post("/django/api/excel/assistant", {"file": self.file}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_assistant_garbage_file(self):
        df = pd.DataFrame({"abc": [1, 2], "xyz": ["a", "b"]})
        f = SimpleUploadedFile("lixo.xlsx", make_xlsx(df),
                               content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r = self.client.post("/django/api/excel/assistant",
                             {"file": f, "question": "o que tens aqui?"},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIn("mapeamento", r.json()["answer"].lower())

    def test_apply_endpoint(self):
        plan = [{"op": "sort", "by": "amount", "ascending": False}]
        r = self.client.post("/django/api/excel/assistant/apply",
                             {"file": self.file, "plan": json.dumps(plan)},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("Sumário", wb.sheetnames)
        self.assertIn("Dados", wb.sheetnames)
        self.assertIn("Fórmulas", wb.sheetnames)
        ws = wb["Dados"]
        self.assertEqual(ws["A1"].value, "Documento")     # cabeçalho original
        self.assertEqual(ws["B2"].value, "Beta")          # 25000 primeiro
        tables = ws.tables
        self.assertIn("tblDados", tables)
        # KPI vivo no Sumário
        self.assertTrue(str(wb["Sumário"]["B5"].value).startswith("=SUM("))
        # folha de fórmulas tem conteúdo
        self.assertTrue(wb["Fórmulas"]["A5"].value)

    def test_apply_invalid_plan_still_works(self):
        r = self.client.post("/django/api/excel/assistant/apply",
                             {"file": self.file, "plan": json.dumps([{"op": "hack"}])},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)  # plano vazio → workbook puro

    def test_viewer_cannot_use_copilot(self):
        viewer = "tok-viewer-cop"
        ApiToken.objects.create(token=viewer, user_id="viewer_cop", role="viewer")
        c = APIClient()
        c.credentials(HTTP_AUTHORIZATION="Bearer " + viewer)
        r = c.post("/django/api/excel/assistant",
                   {"file": self.file, "question": "total?"}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

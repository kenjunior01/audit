"""Testes da Análise Comparativa — o que mudou entre dois períodos (A → B).

Cobre:
- auto_compare: shape, deltas de linhas/total, maiores variações,
  fornecedores novos, padrões (weekend/duplicados/materialidade),
  documentos repetidos ENTRE ficheiros (com flag de mesmo valor)
- veredicto determinístico: estável (ok), moderado (warn), acentuado (high)
- build_compare_report: workbook premium (PK + folhas esperadas)
- API: 200 multipart, 400 sem file_b, relatório xlsx, 403 viewer, 401 sem token
"""
import io

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from dashboard import excel_auto_service as auto
from dashboard.models import ApiToken

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PERIOD_A = pd.DataFrame([
    {"Documento": f"NF-A{i}", "Fornecedor": f"Forn-{i % 3}", "Valor": 100.0 + i * 11.0,
     "Data": f"2025-01-{(i % 28) + 1:02d}", "Categoria": "TI", "Estado": "Pending"}
    for i in range(1, 16)
])

# Período B: mesmos fornecedores base + documento repetido no próprio B
# (NF-REP1 ×2) + documento repetido VS A com valor diferente (NF-A1) +
# fornecedor novo com valor relevante (Novo Ltda).
PERIOD_B = pd.DataFrame([
    {"Documento": f"NF-B{i}", "Fornecedor": f"Forn-{i % 3}", "Valor": 150.0 + i * 11.0,
     "Data": f"2025-02-{(i % 28) + 1:02d}", "Categoria": "TI", "Estado": "Pending"}
    for i in range(1, 16)
] + [
    {"Documento": "NF-REP1", "Fornecedor": "Forn-0", "Valor": 500.0,
     "Data": "2025-02-10", "Categoria": "TI", "Estado": "Pending"},
    {"Documento": "NF-REP1", "Fornecedor": "Forn-0", "Valor": 500.0,
     "Data": "2025-02-11", "Categoria": "TI", "Estado": "Pending"},
    {"Documento": "NF-A1", "Fornecedor": "Forn-1", "Valor": 999.0,
     "Data": "2025-02-12", "Categoria": "TI", "Estado": "Pending"},
    {"Documento": "NF-NOVO1", "Fornecedor": "Novo Ltda", "Valor": 5000.0,
     "Data": "2025-02-20", "Categoria": "Consultoria", "Estado": "Pending"},
])

# Mesmos valores/fornecedores/datas de A, mas IDs diferentes → comparação
# limpa (sem documentos repetidos) para o veredicto "estável".
PERIOD_SAME = PERIOD_A.copy()
PERIOD_SAME["Documento"] = [f"NF-S{i}" for i in range(1, 16)]


def to_xlsx(df, sheet="Dados"):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name=sheet)
    return buf.getvalue()


class CompareServiceTest(TestCase):
    """Análise Comparativa — serviço (deltas, padrões, repetidos, veredicto)."""

    def test_shape_and_deltas(self):
        r = auto.auto_compare(to_xlsx(PERIOD_A), "jan.xlsx",
                              to_xlsx(PERIOD_B), "fev.xlsx")
        for k in ("file_a", "file_b", "overview", "vendors", "patterns",
                  "repeated_ids", "verdict", "level", "message", "insights"):
            self.assertIn(k, r)
        self.assertEqual(r["overview"]["a"]["rows"], 15)
        self.assertEqual(r["overview"]["b"]["rows"], 19)
        self.assertEqual(r["overview"]["delta"]["rows"], 4)
        self.assertGreater(r["overview"]["delta"]["total"], 0)
        # maior aumento = fornecedor novo com +5.000
        top = r["vendors"]["top_increases"][0]
        self.assertEqual(top["vendor"], "Novo Ltda")
        self.assertAlmostEqual(top["delta"], 5000.0, places=2)
        self.assertIn("Novo Ltda", [d["vendor"] for d in r["vendors"]["new"]])
        self.assertIn("weekend", r["patterns"])
        self.assertGreaterEqual(len(r["insights"]), 5)
        self.assertTrue(r["insights"][-1].startswith("Veredicto:"))

    def test_repeated_ids_between_files(self):
        r = auto.auto_compare(to_xlsx(PERIOD_A), "jan.xlsx",
                              to_xlsx(PERIOD_B), "fev.xlsx")
        self.assertEqual(r["repeated_ids"]["count"], 1)   # NF-A1 (REP1 só está em B)
        row = r["repeated_ids"]["rows"][0]
        self.assertEqual(row["id"], "NF-A1")
        self.assertFalse(row["same_amount"])              # 111,00 vs 999,00

    def test_duplicates_pattern_counted(self):
        r = auto.auto_compare(to_xlsx(PERIOD_A), "jan.xlsx",
                              to_xlsx(PERIOD_B), "fev.xlsx")
        self.assertEqual(r["patterns"]["duplicates"]["a"], 0)
        self.assertEqual(r["patterns"]["duplicates"]["b"], 1)  # NF-REP1 ×2

    def test_verdict_stable_when_no_changes(self):
        r = auto.auto_compare(to_xlsx(PERIOD_A), "a.xlsx",
                              to_xlsx(PERIOD_SAME), "b.xlsx")
        self.assertEqual(r["level"], "ok")
        self.assertEqual(r["verdict"], "Períodos estáveis")
        self.assertEqual(r["overview"]["delta"]["total_pct"], 0.0)
        self.assertEqual(r["repeated_ids"]["count"], 0)

    def test_verdict_acentuado_on_big_swing(self):
        big = PERIOD_B.copy()
        big["Valor"] = 5000.0
        r = auto.auto_compare(to_xlsx(PERIOD_A), "a.xlsx", to_xlsx(big), "b.xlsx")
        self.assertEqual(r["level"], "high")
        self.assertEqual(r["verdict"], "Variação acentuada")

    def test_report_is_premium_xlsx(self):
        wb_bytes, fname = auto.build_compare_report(
            to_xlsx(PERIOD_A), "jan.xlsx", to_xlsx(PERIOD_B), "fev.xlsx")
        self.assertTrue(wb_bytes[:2] == b"PK")
        self.assertTrue(fname.startswith("Comparacao_Auto_"))
        wb = load_workbook(io.BytesIO(wb_bytes), read_only=True)
        self.assertIn("Sumário", wb.sheetnames)
        self.assertIn("Documentos repetidos", wb.sheetnames)
        self.assertIn("Padrões", wb.sheetnames)
        self.assertIn("Novos fornecedores", wb.sheetnames)


class CompareApiTest(TestCase):
    """Análise Comparativa — API (multipart, erros acionáveis, permissões)."""

    @classmethod
    def setUpTestData(cls):
        ApiToken.objects.create(token="tok-cmp-viewer", user_id="cmp_viewer", role="viewer")
        ApiToken.objects.create(token="tok-cmp-auditor", user_id="cmp_auditor", role="auditor")

    def setUp(self):
        self.client = APIClient()

    def _fa(self):
        return SimpleUploadedFile("jan.xlsx", to_xlsx(PERIOD_A), content_type=MIME)

    def _fb(self):
        return SimpleUploadedFile("fev.xlsx", to_xlsx(PERIOD_B), content_type=MIME)

    def test_compare_ok(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-cmp-auditor")
        r = self.client.post("/django/api/excel/auto/compare",
                             {"file_a": self._fa(), "file_b": self._fb()},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["overview"]["b"]["rows"], 19)
        self.assertEqual(d["repeated_ids"]["count"], 1)
        self.assertIn("insights", d)

    def test_missing_file_b_actionable(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-cmp-auditor")
        r = self.client.post("/django/api/excel/auto/compare",
                             {"file_a": self._fa()}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("file_b", r.json()["error"])

    def test_report_download(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-cmp-auditor")
        r = self.client.post("/django/api/excel/auto/compare/report",
                             {"file_a": self._fa(), "file_b": self._fb()},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.content[:2], b"PK")
        self.assertIn("Comparacao_Auto", r["Content-Disposition"])

    def test_viewer_forbidden(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-cmp-viewer")
        r = self.client.post("/django/api/excel/auto/compare",
                             {"file_a": self._fa(), "file_b": self._fb()},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_requires_auth(self):
        r = self.client.post("/django/api/excel/auto/compare",
                             {"file_a": self._fa(), "file_b": self._fb()},
                             format="multipart")
        self.assertIn(r.status_code, (status.HTTP_401_UNAUTHORIZED,
                                      status.HTTP_403_FORBIDDEN))

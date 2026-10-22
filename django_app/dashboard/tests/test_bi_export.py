"""Testes da exportação BI → Excel formatado (endpoint + conteúdo xlsx)."""
import io
import uuid

import openpyxl
from django.test import TestCase
from rest_framework.test import APIClient

from dashboard.models import ApiToken

MEAS = {'op': 'agg', 'agg': 'sum', 'field': 'amount'}


class ExportEndpointTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tok = str(uuid.uuid4())
        ApiToken.objects.create(token=cls.tok, user_id="exp@x.pt", role="viewer")

    def _post(self, payload):
        c = APIClient()
        c.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok)
        return c.post("/django/api/bi/export/xlsx", payload, format="json")

    def test_export_group(self):
        payload = {'kind': 'group', 'title': 'Total por fornecedor',
                   'data': {'dimension': 'vendor', 'mode': 'plain',
                            'rows': [{'vendor': 'A', 'value': 100},
                                     {'vendor': 'B', 'value': 250.5}]}}
        r = self._post(payload)
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('.xlsx', r['Content-Disposition'])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Dados']
        self.assertEqual(ws.cell(row=4, column=1).value, 'Membro')
        self.assertEqual(ws.cell(row=5, column=1).value, 'A')
        self.assertEqual(ws.cell(row=6, column=2).value, 250.5)

    def test_export_matrix_with_totals(self):
        payload = {'kind': 'matrix', 'title': 'Matrix',
                   'data': {'row_dim': 'vendor', 'col_dim': 'month',
                            'columns': ['2026-09', '2026-10'],
                            'rows': [{'row': 'A', '2026-09': 10,
                                      '2026-10': 20, 'total': 30}],
                            'col_totals': {'2026-09': 10, '2026-10': 20},
                            'grand_total': 30}}
        r = self._post(payload)
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Dados']
        self.assertEqual(ws.cell(row=5, column=1).value, 'A')
        self.assertEqual(ws.cell(row=5, column=4).value, 30)
        self.assertEqual(ws.cell(row=6, column=4).value, 30)  # total geral

    def test_export_kpi_pareto_series(self):
        base = {'title': 't'}
        r = self._post({**base, 'kind': 'kpi',
                        'data': {'value': 900, 'target': 1000,
                                 'pct_of_target': 0.9, 'status': 'warn',
                                 'trend_30d': -0.1, 'rows': 3}})
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn('KPI', wb.sheetnames)
        r = self._post({**base, 'kind': 'pareto',
                        'data': {'dimension': 'vendor',
                                 'rows': [{'vendor': 'A', 'value': 9,
                                           'share': 0.9, 'cum_share': 0.9,
                                           'vital_few': True}]}})
        self.assertEqual(r.status_code, 200)
        r = self._post({**base, 'kind': 'series',
                        'data': {'granularity': 'month',
                                 'rows': [{'key': '2026-09', 'value': 10,
                                           'moving_avg': 10,
                                           'cumulative': 10}]}})
        self.assertEqual(r.status_code, 200)
        r = self._post({**base, 'kind': 'forecast',
                        'data': {'granularity': 'month', 'slope': 1.0,
                                 'history': [{'key': '2026-09', 'value': 10}],
                                 'forecast': [{'key': '2026-10', 'value': 11,
                                               'lower': 9, 'upper': 13}]}})
        self.assertEqual(r.status_code, 200)

    def test_export_drillthrough_rows(self):
        payload = {'kind': 'drillthrough', 'title': 'Detalhe',
                   'data': {'count': 1, 'rows': [
                       {'transaction_id': 'T1', 'vendor': 'A', 'amount': 10,
                        'currency': 'BRL', 'category': 'Compras',
                        'status': 'Completed', 'user_id': 'u1',
                        'timestamp': '2026-10-15T00:00:00Z'}]}}
        r = self._post(payload)
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Dados']
        self.assertEqual(ws.cell(row=4, column=1).value, 'transaction_id')

    def test_export_requires_data(self):
        r = self._post({'kind': 'group'})
        self.assertEqual(r.status_code, 400)

    def test_export_caps_rows(self):
        rows = [{'vendor': f'V{i}', 'value': i} for i in range(6000)]
        r = self._post({'kind': 'group', 'data': {'dimension': 'vendor',
                                                  'rows': rows}})
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Dados']
        self.assertEqual(ws.max_row, 4 + 5000)  # 4 header + 5000 linhas

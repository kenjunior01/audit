"""Testes das correções de robustez do Excel Studio (fix "não funciona").

Cobre:
- Retry resiliente de folhas no Copiloto (chat + apply): folha errada à
  frente não bloqueia mais o utilizador
- Erro acionável (400) quando o ficheiro é ilegível, com colunas detetadas
- Aliases de rotas que estavam 404: /context/profile/current e
  /settings/integration/current
"""
import io

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from dashboard.models import ApiToken

VALID = pd.DataFrame([
    {"Documento": "NF-1", "Fornecedor": "Alpha", "Valor": 1200.0,
     "Data": "2025-03-04", "Categoria": "TI", "Estado": "Pending"},
    {"Documento": "NF-2", "Fornecedor": "Beta", "Valor": 25000.0,
     "Data": "2025-03-05", "Categoria": "Admin", "Estado": "Approved"},
    {"Documento": "NF-3", "Fornecedor": "Alpha", "Valor": 900.5,
     "Data": "2025-04-10", "Categoria": "TI", "Estado": "Pending"},
])

GARBAGE = pd.DataFrame({"abc": [1, 2], "xyz": ["a", "b"]})


def make_multisheet_xlsx() -> bytes:
    """Folha 1 'Capa' ilegível (capa decorativa), folha 2 'Movimentos' válida."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        GARBAGE.to_excel(w, index=False, sheet_name="Capa")
        VALID.to_excel(w, index=False, sheet_name="Movimentos")
    return buf.getvalue()


class MultiSheetRetryTest(TestCase):
    """A folha escolhida automaticamente não tem dados → tenta as outras."""

    @classmethod
    def setUpTestData(cls):
        cls.token = "tok-fix-1"
        ApiToken.objects.create(token=cls.token, user_id="fix_admin", role="admin")
        cls.bytes_xlsx = make_multisheet_xlsx()

    def setUp(self):
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.token)

    def _file(self):
        return SimpleUploadedFile(
            "relatorio.xlsx", self.bytes_xlsx,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_chat_reads_second_sheet(self):
        r = self.client.post("/django/api/excel/assistant",
                             {"file": self._file(), "question": "total por fornecedor"},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["sheet"], "Movimentos")      # saltou a Capa
        self.assertGreater(d["result"]["total_rows"], 0)

    def test_apply_reads_second_sheet(self):
        r = self.client.post("/django/api/excel/assistant/apply",
                             {"file": self._file(),
                              "plan": '[{"op": "sort", "by": "amount", "ascending": false}]'},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("Dados", wb.sheetnames)
        self.assertEqual(wb["Dados"]["B2"].value, "Beta")  # 25000 primeiro

    def test_unreadable_apply_gives_actionable_400(self):
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            GARBAGE.to_excel(w, index=False, sheet_name="Unica")
        f = SimpleUploadedFile("lixo.xlsx", buf.getvalue(),
                               content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r = self.client.post("/django/api/excel/assistant/apply",
                             {"file": f, "plan": "[]"}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        msg = r.json().get("error", "")
        self.assertIn("abc", msg)                        # colunas detetadas
        self.assertIn("Importar", msg)                   # caminho de resolução

    def test_unreadable_chat_is_friendly_200(self):
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            GARBAGE.to_excel(w, index=False, sheet_name="Unica")
        f = SimpleUploadedFile("lixo.xlsx", buf.getvalue(),
                               content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r = self.client.post("/django/api/excel/assistant",
                             {"file": f, "question": "total?"}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["result"]["total_rows"], 0)
        self.assertIn("abc", d["answer"])                # agora lista as colunas

    def test_explicit_sheet_is_respected(self):
        r = self.client.post("/django/api/excel/assistant",
                             {"file": self._file(), "question": "total?",
                              "sheet": "Movimentos"}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.json()["sheet"], "Movimentos")


class RouteAliasTest(TestCase):
    """Rotas que o frontend chama e estavam 404."""

    @classmethod
    def setUpTestData(cls):
        cls.token = "tok-fix-2"
        ApiToken.objects.create(token=cls.token, user_id="fix_admin2", role="admin")

    def setUp(self):
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.token)

    def test_context_profile_current(self):
        r = self.client.get("/django/api/context/profile/current")
        self.assertEqual(r.status_code, status.HTTP_200_OK)

    def test_settings_integration_current(self):
        r = self.client.get("/django/api/settings/integration/current")
        self.assertEqual(r.status_code, status.HTTP_200_OK)

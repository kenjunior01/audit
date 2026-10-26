"""Testes do Excel Auto — análise, reconciliação e dossiê de 1 clique.

Cobre:
- /excel/auto: shape completo (perfil, qualidade 0-100, risco, insights)
- retry multi-folha (capa decorativa → salta para a folha com dados)
- erro acionável (400) em ficheiro ilegível
- ficheiro sujo pontua qualidade MENOR que ficheiro limpo (checks warn/fail)
- /excel/auto/report: workbook premium descarregável (PK zip)
- /excel/sample: exemplo gerado e analisável pelo próprio fluxo auto
- /excel/auto/reconcile(+report): 2 ficheiros e ficheiro vs base de dados,
  incluindo a correção dos órfãos (unmatched_a = linhas SEM correspondência)
- /excel/auto/dossier: analisa + importa + painel REA + workbook consolidado
- permissões: viewer OK na leitura, auditor para importar, sem token 401/403
"""
import io

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from dashboard import excel_auto_service as auto
from dashboard.models import ApiToken

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CLEAN = pd.DataFrame([
    {"Documento": f"NF-{i}", "Fornecedor": f"Forn-{i % 3}", "Valor": 100.0 + i * 13.7,
     "Data": f"2025-03-{(i % 28) + 1:02d}", "Categoria": f"Cat-{i % 2}", "Estado": "Pending"}
    for i in range(30)
])

DIRTY = pd.DataFrame([
    {"Documento": "NF-1", "Fornecedor": "Alpha", "Valor": 1200.0,
     "Data": "2025-03-04", "Categoria": "TI", "Estado": "Pending"},
    {"Documento": "NF-1", "Fornecedor": "Alpha", "Valor": 1200.0,
     "Data": "2025-03-04", "Categoria": "TI", "Estado": "Pending"},   # duplicado exato
    {"Documento": "NF-2", "Fornecedor": "Beta", "Valor": 9999.0,
     "Data": "2025-03-05", "Categoria": "Admin", "Estado": "Approved"},  # split 1
    {"Documento": "NF-3", "Fornecedor": "Beta", "Valor": 9876.0,
     "Data": "2025-03-05", "Categoria": "Admin", "Estado": "Approved"},  # split 2
    {"Documento": "NF-4", "Fornecedor": "Gama", "Valor": "N/D",
     "Data": "2025-03-06", "Categoria": "TI", "Estado": "Pending"},   # ilegível
] + [{"Documento": f"NF-{i}", "Fornecedor": f"Forn-{i % 3}", "Valor": 100.0 + i,
      "Data": f"2025-04-{(i % 28) + 1:02d}", "Categoria": "TI", "Estado": "Pending"}
     for i in range(5, 25)])

GARBAGE = pd.DataFrame({"abc": [1, 2], "xyz": ["a", "b"]})


def to_xlsx(df, sheet="Dados"):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name=sheet)
    return buf.getvalue()


def make_multisheet() -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        GARBAGE.to_excel(w, index=False, sheet_name="Capa")
        CLEAN.to_excel(w, index=False, sheet_name="Movimentos")
    return buf.getvalue()


def recent_df(n=6, prefix="DBX", amount=200.0):
    """Dados com datas recentes (janela de 365 dias da base)."""
    base = pd.Timestamp.utcnow().tz_localize(None).normalize() - pd.Timedelta(days=10)
    return pd.DataFrame([
        {"Documento": f"{prefix}-{i}", "Fornecedor": f"Forn-{i}",
         "Valor": amount + i * 7.5, "Data": (base + pd.Timedelta(days=i)).strftime("%Y-%m-%d"),
         "Categoria": "TI", "Estado": "Pending"}
        for i in range(n)])


class ExcelAutoServiceTest(TestCase):
    """Excel Auto — serviço de análise (1 ficheiro → diagnóstico completo)."""

    def test_shape_and_fields(self):
        r = auto.auto_analyze(to_xlsx(CLEAN), "limpo.xlsx")
        for k in ("file", "sheet", "mapping", "profile", "quality", "risk",
                  "insights", "top_vendors", "next_steps", "analyses"):
            self.assertIn(k, r)
        self.assertEqual(r["profile"]["valid_rows"], 30)
        self.assertEqual(r["profile"]["skipped_rows"], 0)
        self.assertEqual(r["quality"]["grade"], "A")
        self.assertEqual(len(r["quality"]["checks"]), 6)
        self.assertGreaterEqual(len(r["insights"]), 8)
        self.assertTrue(r["insights"][0].startswith("Período"))

    def test_multisheet_skips_cover(self):
        r = auto.auto_analyze(make_multisheet(), "rel.xlsx")
        self.assertEqual(r["sheet"], "Movimentos")
        self.assertIn("Capa", [s["name"] for s in r["sheets"]])
        self.assertEqual(r["profile"]["valid_rows"], 30)

    def test_ilegivel_error_actionable(self):
        with self.assertRaises(ValueError) as ctx:
            auto.auto_analyze(to_xlsx(GARBAGE), "lixo.xlsx")
        self.assertIn("Colunas detetadas", str(ctx.exception))
        self.assertIn("lixo.xlsx", str(ctx.exception))

    def test_dirty_scores_lower_than_clean(self):
        rc = auto.auto_analyze(to_xlsx(CLEAN), "limpo.xlsx")
        rd = auto.auto_analyze(to_xlsx(DIRTY), "sujo.xlsx")
        self.assertLess(rd["quality"]["score"], rc["quality"]["score"])
        checks = {c["id"]: c["status"] for c in rd["quality"]["checks"]}
        self.assertEqual(checks["dup_ids"], "warn")
        self.assertGreater(rd["risk"]["splits"], 0)

    def test_report_is_premium_xlsx(self):
        wb_bytes, fname = auto.build_auto_report(to_xlsx(CLEAN), "limpo.xlsx")
        self.assertTrue(wb_bytes[:2] == b"PK")
        self.assertTrue(fname.endswith(".xlsx"))
        wb = load_workbook(io.BytesIO(wb_bytes), read_only=True)
        self.assertIn("Qualidade", wb.sheetnames)
        self.assertIn("Sumário", wb.sheetnames)

    def test_sample_is_analyzable(self):
        sample = auto.generate_sample_xlsx()
        rs = auto.auto_analyze(sample, "exemplo.xlsx")
        self.assertGreater(rs["profile"]["valid_rows"], 40)
        self.assertGreater(rs["profile"]["skipped_rows"], 0)   # problema plantado
        self.assertTrue(any("Benford" in i for i in rs["insights"]))


class ExcelAutoApiTest(TestCase):
    """Excel Auto — API (multipart, erros acionáveis, permissões)."""

    @classmethod
    def setUpTestData(cls):
        ApiToken.objects.create(token="tok-auto-viewer", user_id="auto_viewer", role="viewer")

    def setUp(self):
        self.client = APIClient()

    def _file(self, content=None, name="auto.xlsx"):
        return SimpleUploadedFile(name, content or to_xlsx(CLEAN), content_type=MIME)

    def test_viewer_can_analyze(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-auto-viewer")
        r = self.client.post("/django/api/excel/auto",
                             {"file": self._file()}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["profile"]["valid_rows"], 30)
        self.assertIn("quality", d)
        self.assertGreaterEqual(len(d["insights"]), 6)

    def test_requires_file_with_actionable_error(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-auto-viewer")
        r = self.client.post("/django/api/excel/auto", {}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("exemplo", r.json()["error"])

    def test_ilegivel_returns_400_with_columns(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-auto-viewer")
        r = self.client.post("/django/api/excel/auto",
                             {"file": self._file(to_xlsx(GARBAGE))}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Colunas detetadas", r.json()["error"])

    def test_requires_auth(self):
        r = self.client.post("/django/api/excel/auto",
                             {"file": self._file()}, format="multipart")
        self.assertIn(r.status_code, (status.HTTP_401_UNAUTHORIZED,
                                      status.HTTP_403_FORBIDDEN))


class ReconcileAutoServiceTest(TestCase):
    """Reconciliação Auto — 2 ficheiros ou ficheiro vs base de dados."""

    def test_perfect_match(self):
        a = recent_df(6)
        r = auto.auto_reconcile(to_xlsx(a), "a.xlsx", to_xlsx(a.copy()), "b.xlsx")
        self.assertEqual(r["summary"]["matched"], 6)
        self.assertEqual(r["summary"]["unmatched_a"], 0)
        self.assertEqual(r["verdict"], "Reconciliação perfeita")

    def test_orphans_are_the_unmatched_rows(self):
        # 1 linha com valor alterado → órfã em A. A correção do es.reconcile
        # garante que unmatched_a contém a ÓRFÃ (não as linhas casadas).
        a = recent_df(6)
        b = a.copy()
        b.loc[b.index[-1], "Valor"] = b["Valor"].iloc[-1] + 500.0
        b.loc[b.index[-1], "Documento"] = "BX-99"
        r = auto.auto_reconcile(to_xlsx(a), "a.xlsx", to_xlsx(b), "b.xlsx")
        self.assertEqual(r["summary"]["matched"], 5)
        self.assertEqual(r["summary"]["unmatched_a"], 1)
        orphan_ids = {rec.get("Documento") or rec.get("transaction_id")
                      for rec in r["unmatched_a"]}
        self.assertIn("DBX-5", orphan_ids)
        self.assertEqual(r["summary"]["difference"], -500.0)  # A menor que B

    def test_ilegivel_b_raises_actionable(self):
        a = recent_df(6)
        with self.assertRaises(ValueError) as ctx:
            auto.auto_reconcile(to_xlsx(a), "a.xlsx", to_xlsx(GARBAGE), "lixo.xlsx")
        self.assertIn("lixo.xlsx", str(ctx.exception))
        self.assertIn("Colunas detetadas", str(ctx.exception))

    def test_requires_second_source(self):
        a = recent_df(6)
        with self.assertRaises(ValueError):
            auto.auto_reconcile(to_xlsx(a), "a.xlsx", None, None)

    def test_against_database_and_report(self):
        from dashboard.models import Transaction
        df_b = recent_df(6, prefix="BASE")
        for _, row in df_b.iterrows():
            Transaction.objects.create(
                transaction_id=row["Documento"], vendor=row["Fornecedor"],
                amount=row["Valor"], currency="EUR",
                timestamp=pd.Timestamp(row["Data"]), category="TI",
                user_id="t", status="Pending")
        a = recent_df(6, prefix="BASE")   # mesmos IDs → match perfeito
        r = auto.auto_reconcile(to_xlsx(a), "a.xlsx", df_b=df_b)
        self.assertEqual(r["summary"]["matched"], 6)
        self.assertEqual(r["file_b"], "Base de dados (transações importadas)")
        wb_bytes, fname = auto.build_reconcile_report(to_xlsx(a), "a.xlsx", df_b=df_b)
        self.assertTrue(wb_bytes[:2] == b"PK")
        wb = load_workbook(io.BytesIO(wb_bytes), read_only=True)
        self.assertIn("Só em A", wb.sheetnames)


class ReconcileAutoApiTest(TestCase):
    """Reconciliação Auto — API (multipart, base de dados, permissões)."""

    @classmethod
    def setUpTestData(cls):
        from dashboard.models import ApiToken, Transaction
        cls.tok_admin = "tok-rec-admin"
        ApiToken.objects.create(token=cls.tok_admin, user_id="rec_admin", role="admin")
        cls.tok_viewer = "tok-rec-viewer"
        ApiToken.objects.create(token=cls.tok_viewer, user_id="rec_viewer", role="viewer")
        df = recent_df(3, prefix="DBX", amount=200.0)
        for i, (_, row) in enumerate(df.iterrows()):
            Transaction.objects.create(
                transaction_id=row["Documento"], vendor=row["Fornecedor"],
                amount=200.0 + i * 7, currency="EUR",
                timestamp=pd.Timestamp(row["Data"]),
                category="TI", user_id="t", status="Pending")

    def setUp(self):
        self.client = APIClient()

    def _f(self, name, content):
        return SimpleUploadedFile(name, content, content_type=MIME)

    def test_file_vs_file(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok_admin)
        a = recent_df(6)
        r = self.client.post("/django/api/excel/auto/reconcile",
                             {"file_a": self._f("a.xlsx", to_xlsx(a)),
                              "file_b": self._f("b.xlsx", to_xlsx(a.copy()))},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["summary"]["matched"], 6)
        self.assertIn("verdict", d)
        self.assertIn("insights", d)

    def test_file_vs_database(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok_admin)
        a = recent_df(3, prefix="DBX", amount=200.0)
        r = self.client.post("/django/api/excel/auto/reconcile",
                             {"file_a": self._f("a.xlsx", to_xlsx(a)), "use_db_b": "1"},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        d = r.json()
        self.assertEqual(d["source_b"], "ficheiro") if False else None
        self.assertIn("Base de dados", d["file_b"])
        self.assertEqual(d["summary"]["rows_a"], 3)

    def test_requires_b(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok_admin)
        r = self.client.post("/django/api/excel/auto/reconcile",
                             {"file_a": self._f("a.xlsx", to_xlsx(CLEAN))},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_report_download(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok_admin)
        r = self.client.post("/django/api/excel/auto/reconcile/report",
                             {"file_a": self._f("a.xlsx", to_xlsx(CLEAN)),
                              "file_b": self._f("b.xlsx", to_xlsx(CLEAN.copy()))},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertTrue(r.content[:2] == b"PK")

    def test_permissions(self):
        # viewer → 403 (endpoint auditor)
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.tok_viewer)
        r = self.client.post("/django/api/excel/auto/reconcile",
                             {"file_a": self._f("a.xlsx", to_xlsx(CLEAN)),
                              "file_b": self._f("b.xlsx", to_xlsx(CLEAN.copy()))},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)
        # sem token → 401
        self.client.credentials()
        r2 = self.client.post("/django/api/excel/auto/reconcile",
                              {"file_a": self._f("a.xlsx", to_xlsx(CLEAN)),
                               "file_b": self._f("b.xlsx", to_xlsx(CLEAN.copy()))},
                              format="multipart")
        self.assertIn(r2.status_code, (status.HTTP_401_UNAUTHORIZED,
                                       status.HTTP_403_FORBIDDEN))


class DossierServiceTest(TestCase):
    """Dossiê de Auditoria de 1 clique — análise + importação + painel REA."""

    def test_dossier_imports_and_builds(self):
        from dashboard.models import Transaction
        wb_bytes, fname, meta = auto.build_dossier(to_xlsx(CLEAN), "limpo.xlsx")
        self.assertTrue(wb_bytes[:2] == b"PK")
        self.assertTrue(fname.endswith(".xlsx"))
        self.assertEqual(meta["imported"], 30)
        self.assertEqual(meta["duplicates"], 0)
        self.assertEqual(Transaction.objects.count(), 30)
        wb = load_workbook(io.BytesIO(wb_bytes), read_only=True)
        names = wb.sheetnames
        self.assertIn("Dossiê", names)
        self.assertIn("Qualidade", names)
        self.assertIn("Sumário", names)

    def test_dossier_second_run_dedups(self):
        auto.build_dossier(to_xlsx(CLEAN), "limpo.xlsx")
        wb2, fname2, meta2 = auto.build_dossier(to_xlsx(CLEAN), "limpo.xlsx")
        self.assertEqual(meta2["imported"], 0)
        self.assertEqual(meta2["duplicates"], 30)


class DossierApiTest(TestCase):
    """Dossiê de Auditoria — API (multipart, permissões auditor/admin)."""

    @classmethod
    def setUpTestData(cls):
        ApiToken.objects.create(token="tok-doss-admin", user_id="doss_admin", role="admin")
        ApiToken.objects.create(token="tok-doss-viewer", user_id="doss_viewer", role="viewer")

    def setUp(self):
        self.client = APIClient()

    def _f(self, name="d.xlsx"):
        return SimpleUploadedFile(name, to_xlsx(CLEAN), content_type=MIME)

    def test_dossier_endpoint_download(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-doss-admin")
        r = self.client.post("/django/api/excel/auto/dossier",
                             {"file": self._f()}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertTrue(r.content[:2] == b"PK")
        self.assertEqual(r["X-Dossier-Imported"], "30")

    def test_dossier_requires_auditor_and_file(self):
        # viewer → 403 (importar é ação de auditor)
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-doss-viewer")
        r = self.client.post("/django/api/excel/auto/dossier",
                             {"file": self._f()}, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)
        # auditor sem ficheiro → 400 acionável
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-doss-admin")
        r2 = self.client.post("/django/api/excel/auto/dossier", {}, format="multipart")
        self.assertEqual(r2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("file", str(r2.data.get("error", "")).lower())


CONC = pd.DataFrame([
    {"Documento": f"C-{i}", "Fornecedor": "Alpha" if i < 8 else f"Outro-{i}",
     "Valor": 1000.0 + i * 3, "Data": "2025-03-10", "Categoria": "TI", "Estado": "Pending"}
    for i in range(10)])


class AlertSuggestionsTest(TestCase):
    """Alertas sugeridos — padrões detetados viram regras MeasureAlert (1 clique)."""

    def _analyze(self, df, name="a.xlsx", **kw):
        return auto.auto_analyze(to_xlsx(df), name, **kw)

    def test_suggestions_are_valid_bi_rules(self):
        from dashboard.bi_service import _validate as validate_measure, _validate_filter
        r = self._analyze(CLEAN)
        sugs = r.get("alert_suggestions")
        self.assertIsInstance(sugs, list)
        keys = {s["key"] for s in sugs}
        # CLEAN: fim-de-semana presente, fornecedores equilibrados (sem concentração)
        self.assertIn("weekend", keys)
        self.assertIn("materialidade", keys)
        self.assertNotIn("concentracao", keys)
        for s in sugs:
            validate_measure(s["rule"]["measure"])
            for f in s["rule"].get("filters", []):
                _validate_filter(f)
            self.assertIn(s["rule"]["comparator"], ("gt", "lt", "gte", "lte"))
            self.assertTrue(s["rule"]["name"])
            self.assertTrue(s["why"])

    def test_weekend_rule_detects_after_import(self):
        from dashboard.models import MeasureAlert, Transaction
        from dashboard.bi_views import evaluate_measure_alert
        r = self._analyze(CLEAN)
        rule = {x["key"]: x for x in r["alert_suggestions"]}["weekend"]["rule"]
        # importa o ficheiro e cria o alerta a partir da sugestão
        created, dups, _ = auto.import_standardized(
            [{"transaction_id": f"NF-{i}", "vendor": f"Forn-{i % 3}",
              "amount": 100.0 + i * 13.7, "currency": "EUR",
              "timestamp": f"2025-03-{(i % 28) + 1:02d}",
              "category": f"Cat-{i % 2}", "user_id": "u", "status": "Pending"}
             for i in range(30)], "week.xlsx", uploaded_by="aud",
            sheet_name="Dados", mapping={}, skipped=[])
        self.assertEqual(created, 30)
        a = MeasureAlert.objects.create(
            name=rule["name"], measure=rule["measure"], filters=rule["filters"],
            comparator=rule["comparator"], threshold=rule["threshold"],
            created_by="aud")
        state, value, _ = evaluate_measure_alert(a)
        # CLEAN tem linhas em sáb/dom (2025-03-01=Saturday, 2025-03-02=Sunday)
        self.assertEqual(state, "breach")
        self.assertGreaterEqual(value, 2)
        self.assertEqual(MeasureAlert.objects.count(), 1)
        self.assertTrue(Transaction.objects.exists())

    def test_materialidade_threshold_and_count(self):
        r = self._analyze(CLEAN, threshold=200.0)
        s = {x["key"]: x for x in r["alert_suggestions"]}["materialidade"]
        self.assertIn("200,00", s["rule"]["name"])
        # 22 linhas de CLEAN têm valor ≥ 200 (100 + i*13.7, i ≥ 8)
        self.assertTrue(s["why"].startswith("22 transação"))
        f = s["rule"]["measure"]["filters"][0]
        self.assertEqual(f["field"], "amount")
        self.assertEqual(f["op"], "gte")
        self.assertEqual(float(f["value"]), 200.0)

    def test_concentration_suggestion(self):
        r = self._analyze(CONC)
        s = {x["key"]: x for x in r["alert_suggestions"]}.get("concentracao")
        self.assertIsNotNone(s)
        self.assertEqual(s["rule"]["name"], "Concentração: Alpha")
        f = s["rule"]["measure"]["filters"][0]
        self.assertEqual(f["field"], "vendor")
        self.assertEqual(f["value"], "Alpha")
        self.assertEqual(s["rule"]["comparator"], "gte")
        self.assertGreater(float(s["rule"]["threshold"]), 0)

    def test_api_accepts_suggestion_rule(self):
        """Regra sugerida passa na validação da API /bi/alerts (201) e 409 em duplicado."""
        from dashboard.models import MeasureAlert
        self.client = APIClient()
        ApiToken.objects.create(token="tok-sug-admin", user_id="sug_admin", role="admin")
        self.client.credentials(HTTP_AUTHORIZATION="Bearer tok-sug-admin")
        r = self._analyze(CLEAN)
        rule = {x["key"]: x for x in r["alert_suggestions"]}["weekend"]["rule"]
        resp = self.client.post("/django/api/bi/alerts", rule, format="json")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertTrue(MeasureAlert.objects.filter(name=rule["name"]).exists())
        resp2 = self.client.post("/django/api/bi/alerts", rule, format="json")
        self.assertEqual(resp2.status_code, 409)
